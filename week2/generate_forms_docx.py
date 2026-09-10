import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = docx.Document()

# Set standard margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Title
title_p = doc.add_paragraph()
title_run = title_p.add_run("Unit 6: Microcontroller Systems — Week 2 Assessment Form")
title_run.font.name = "Arial"
title_run.font.size = Pt(18)
title_run.font.bold = True
title_run.font.color.rgb = RGBColor(0, 151, 157) # Arduino Teal
title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

subtitle_p = doc.add_paragraph()
sub_run = subtitle_p.add_run("From PC to Embedded Controller (BTEC Criteria A1, B1)")
sub_run.font.name = "Arial"
sub_run.font.size = Pt(12)
sub_run.font.italic = True
subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_paragraph() # Spacing

# Section 1: Student Information
h1 = doc.add_heading("Section 1: Student Information", level=1)
doc.add_paragraph("1. Full Name:\n__________________________________________________")
doc.add_paragraph("2. Student ID / Group Cohort:\n__________________________________________________")

# Section 2: Diagnostic Engineering Toolbox
doc.add_heading("Section 2: Diagnostic Engineering Toolbox", level=1)
doc.add_paragraph("3. What is the fundamental difference between a digital and an analogue signal?")
doc.add_paragraph("A. Digital has discrete 2-state logic (0V/5V), while analogue varies continuously.")
doc.add_paragraph("B. Analogue signals can only have 0V or 5V values.")
doc.add_paragraph("C. Digital signals are only found in power plants.")

doc.add_paragraph("4. Why does an LED require a current-limiting resistor (e.g. 220 Ohm)?")
doc.add_paragraph("A. To prevent excessive current from damaging the LED and microcontroller pin.")
doc.add_paragraph("B. To step up 5V to 12V.")
doc.add_paragraph("C. To convert direct current into alternating current.")

doc.add_paragraph("5. In the Input -> Processing -> Output control model, what role does a push button play?")
doc.add_paragraph("A. An Input sensor providing binary voltage state (HIGH/LOW).")
doc.add_paragraph("B. A Processing CPU unit executing instructions.")
doc.add_paragraph("C. An Output actuator creating light or mechanical motion.")

doc.add_paragraph("6. What is the maximum recommended continuous DC current output from an Arduino Uno I/O pin?")
doc.add_paragraph("A. 20 mA (40 mA absolute max rating)")
doc.add_paragraph("B. 500 mA")
doc.add_paragraph("C. 2.5 A")

# Section 3: PC vs Embedded Microcontroller Architecture
doc.add_heading("Section 3: PC vs Microcontroller Architecture", level=1)
doc.add_paragraph("7. Compare the role of an Operating System (PC) with Bare-Metal Firmware (Arduino):")
doc.add_paragraph("________________________________________________________________________________________________________________________")

doc.add_paragraph("8. State the primary purpose of a microcontroller in industrial engineering:")
doc.add_paragraph("A. Special-purpose computer designed to control physical hardware and real-time processes directly.")
doc.add_paragraph("B. General-purpose computer built for multi-monitor video editing and high-volume data streaming.")
doc.add_paragraph("C. A device used solely for storage backup.")

# Section 4: Arduino Uno Rev3 Datasheet Scavenger Hunt
doc.add_heading("Section 4: Arduino Uno Rev3 Datasheet Scavenger Hunt", level=1)
doc.add_paragraph("9. What is the total internal Flash program memory size on the ATmega328P?")
doc.add_paragraph("A. 32 KB (0.5 KB used by bootloader)")
doc.add_paragraph("B. 512 MB")
doc.add_paragraph("C. 2 KB")

doc.add_paragraph("10. What is the internal SRAM working memory size on the ATmega328P?")
doc.add_paragraph("A. 2 KB")
doc.add_paragraph("B. 32 KB")
doc.add_paragraph("C. 16 MB")

doc.add_paragraph("11. What is the recommended external DC input voltage range via the Barrel Jack / VIN?")
doc.add_paragraph("A. 7 V to 12 V DC")
doc.add_paragraph("B. 1.5 V to 3.3 V DC")
doc.add_paragraph("C. 24 V to 48 V DC")

doc.add_paragraph("12. What clock frequency is provided by the external crystal oscillator?")
doc.add_paragraph("A. 16 MHz")
doc.add_paragraph("B. 3.2 GHz")
doc.add_paragraph("C. 100 kHz")

# Section 5: Practical Programming & Circuit Design Evidence
doc.add_heading("Section 5: Practical Programming & Circuit Design Evidence", level=1)
doc.add_paragraph("13. Explain the execution difference between void setup() and void loop():")
doc.add_paragraph("________________________________________________________________________________________________________________________")

doc.add_paragraph("14. Why is INPUT_PULLUP used when connecting a pushbutton switch?")
doc.add_paragraph("A. It activates the internal 20k-50k Ohm pull-up resistor to 5V, preventing floating input states.")
doc.add_paragraph("B. It doubles the clock speed of the processor.")
doc.add_paragraph("C. It converts the button into an analogue potentiometer.")

doc.add_paragraph("15. Paste your verified C++ sketch for the Dual-Button & Dual-LED Start/Stop Station:")
doc.add_paragraph("________________________________________________________________________________________________________________________\n________________________________________________________________________________________________________________________\n________________________________________________________________________________________________________________________")

# Image Upload Deliverables
doc.add_paragraph("16. [File Upload] Upload screenshot of your Tinkercad Circuit Simulation (Evidence of schematic wiring and simulation run):")
doc.add_paragraph("[ File Upload Item: Student attaches .png / .jpg / .pdf ]")

doc.add_paragraph("17. [File Upload] Upload clear photograph of your Physical Breadboard & Arduino Build (Evidence of hardware assembly and operation):")
doc.add_paragraph("[ File Upload Item: Student attaches .png / .jpg / .pdf ]")

doc.save("/home/tayo/projects/work/unit6_v2/week2/forms.docx")
print("forms.docx updated successfully with image upload items!")
